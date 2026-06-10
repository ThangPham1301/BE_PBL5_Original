from datetime import date
from io import BytesIO

from django.contrib.auth import get_user_model
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from apps.attendance.models import AttendanceLog
from apps.employees.models import Department, Employee


class ReportExportTests(APITestCase):
    def setUp(self):
        user_model = get_user_model()
        self.department = Department.objects.create(name='Kỹ thuật')
        other_department = Department.objects.create(name='Kế toán')

        self.admin = user_model.objects.create_user(
            username='report_admin',
            password='test-password',
            role='admin',
        )
        manager_user = user_model.objects.create_user(
            username='report_manager',
            password='test-password',
            role='manager',
        )
        self.manager = Employee.objects.create(
            user=manager_user,
            employee_id='MGR-001',
            department=self.department,
            date_joined=date(2026, 1, 1),
        )
        employee_user = user_model.objects.create_user(
            username='report_employee',
            first_name='An',
            last_name='Nguyễn',
            password='test-password',
            role='employee',
        )
        self.employee = Employee.objects.create(
            user=employee_user,
            employee_id='EMP-001',
            department=self.department,
            date_joined=date(2026, 1, 1),
        )
        other_user = user_model.objects.create_user(
            username='other_employee',
            password='test-password',
            role='employee',
        )
        self.other_employee = Employee.objects.create(
            user=other_user,
            employee_id='EMP-002',
            department=other_department,
            date_joined=date(2026, 1, 1),
        )

        AttendanceLog.objects.create(
            employee=self.employee,
            date=date(2026, 6, 2),
            status=AttendanceLog.Status.PRESENT,
        )
        AttendanceLog.objects.create(
            employee=self.employee,
            date=date(2026, 6, 3),
            status=AttendanceLog.Status.LATE,
        )

    def test_admin_can_export_a_valid_excel_workbook(self):
        self.client.force_authenticate(user=self.admin)

        response = self.client.get('/api/reports/export/?month=6&year=2026')

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

        worksheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(
            worksheet['A1'].value,
            'BÁO CÁO CHẤM CÔNG THÁNG 6/2026',
        )
        self.assertEqual(
            [worksheet.cell(3, column).value for column in range(1, 10)],
            [
                'STT',
                'Mã NV',
                'Họ tên',
                'Phòng ban',
                'Có mặt',
                'Đi trễ',
                'Vắng',
                'Nghỉ phép',
                'Ngày công',
            ],
        )

        employee_rows = {
            worksheet.cell(row, 2).value: [
                worksheet.cell(row, column).value
                for column in range(5, 10)
            ]
            for row in range(4, worksheet.max_row + 1)
        }
        self.assertEqual(employee_rows['EMP-001'], [1, 1, 0, 0, 2])

    def test_manager_export_is_limited_to_their_department(self):
        self.client.force_authenticate(user=self.manager.user)

        response = self.client.get('/api/reports/export/?month=6&year=2026')

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        worksheet = load_workbook(BytesIO(response.content)).active
        exported_ids = {
            worksheet.cell(row, 2).value
            for row in range(4, worksheet.max_row + 1)
        }
        self.assertEqual(exported_ids, {'MGR-001', 'EMP-001'})

    def test_employee_cannot_export_reports(self):
        self.client.force_authenticate(user=self.employee.user)

        response = self.client.get('/api/reports/export/?month=6&year=2026')

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
