import calendar
from datetime import date
from io import BytesIO

from django.db.models import Count, Q
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from rest_framework import viewsets
from rest_framework.decorators import action
from rest_framework.response import Response

from apps.accounts.permissions import IsAdminOrManager
from apps.employees.models import Employee
from .serializers import ExportQuerySerializer, MonthlyReportQuerySerializer


class ReportViewSet(viewsets.ViewSet):
    permission_classes = [IsAdminOrManager]

    def _get_report_rows(self, user, month, year, department_id=None):
        _, days_in_month = calendar.monthrange(year, month)
        start = date(year, month, 1)
        end = date(year, month, days_in_month)

        employees = Employee.objects.filter(is_active=True).select_related(
            'user',
            'department',
        )
        if user.is_manager:
            try:
                employees = employees.filter(department=user.employee.department)
            except Employee.DoesNotExist:
                employees = employees.none()
        elif department_id:
            employees = employees.filter(department_id=department_id)

        log_filter = Q(attendance_logs__date__range=(start, end))
        employees = employees.annotate(
            present_count=Count(
                'attendance_logs',
                filter=log_filter & Q(attendance_logs__status='present'),
            ),
            late_count=Count(
                'attendance_logs',
                filter=log_filter & Q(attendance_logs__status='late'),
            ),
            absent_count=Count(
                'attendance_logs',
                filter=log_filter & Q(attendance_logs__status='absent'),
            ),
            leave_count=Count(
                'attendance_logs',
                filter=log_filter & Q(attendance_logs__status='leave'),
            ),
        ).order_by('employee_id')

        return [
            {
                'employee_id': employee.employee_id,
                'employee_name': (
                    employee.user.get_full_name() or employee.user.username
                ),
                'department': (
                    employee.department.name if employee.department else None
                ),
                'total_days': days_in_month,
                'present': employee.present_count,
                'late': employee.late_count,
                'absent': employee.absent_count,
                'leave': employee.leave_count,
                'working_days': employee.present_count + employee.late_count,
            }
            for employee in employees
        ]

    @action(detail=False, methods=['get'], url_path='monthly')
    def monthly(self, request):
        serializer = MonthlyReportQuerySerializer(data=request.query_params)
        serializer.is_valid(raise_exception=True)
        params = serializer.validated_data

        report_rows = self._get_report_rows(
            request.user,
            params['month'],
            params['year'],
            params.get('department_id'),
        )

        return Response({
            'success': True,
            'data': {
                'month': params['month'],
                'year': params['year'],
                'employees': report_rows,
            },
            'message': '',
        })

    @action(detail=False, methods=['get'], url_path='export')
    def export(self, request):
        serializer = ExportQuerySerializer(data=request.query_params)
        serializer.is_valid(raise_exception=True)
        params = serializer.validated_data
        month = params['month']
        year = params['year']

        report_rows = self._get_report_rows(
            request.user,
            month,
            year,
            params.get('department_id'),
        )

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = f'Báo cáo T{month}-{year}'

        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin'),
        )
        header_fill = PatternFill(
            start_color='4472C4',
            end_color='4472C4',
            fill_type='solid',
        )
        header_font = Font(bold=True, size=11, color='FFFFFF')

        worksheet.merge_cells('A1:I1')
        worksheet['A1'] = f'BÁO CÁO CHẤM CÔNG THÁNG {month}/{year}'
        worksheet['A1'].font = Font(bold=True, size=14)
        worksheet['A1'].alignment = Alignment(horizontal='center')

        headers = [
            'STT',
            'Mã NV',
            'Họ tên',
            'Phòng ban',
            'Có mặt',
            'Đi trễ',
            'Vắng',
            'Nghỉ phép',
            'Ngày công',
        ]
        for column, header in enumerate(headers, 1):
            cell = worksheet.cell(row=3, column=column, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center')

        for row_number, (index, report_row) in enumerate(
            enumerate(report_rows, 1),
            start=4,
        ):
            values = [
                index,
                report_row['employee_id'],
                report_row['employee_name'],
                report_row['department'] or '',
                report_row['present'],
                report_row['late'],
                report_row['absent'],
                report_row['leave'],
                report_row['working_days'],
            ]
            for column, value in enumerate(values, 1):
                cell = worksheet.cell(
                    row=row_number,
                    column=column,
                    value=value,
                )
                cell.border = thin_border
                if column >= 5:
                    cell.alignment = Alignment(horizontal='center')

        column_widths = [6, 14, 25, 20, 10, 10, 10, 12, 12]
        for column, width in enumerate(column_widths, 1):
            worksheet.column_dimensions[chr(64 + column)].width = width

        worksheet.freeze_panes = 'A4'
        worksheet.auto_filter.ref = f'A3:I{max(3, len(report_rows) + 3)}'

        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer.getvalue(),
            content_type=(
                'application/vnd.openxmlformats-officedocument.'
                'spreadsheetml.sheet'
            ),
        )
        response['Content-Disposition'] = (
            f'attachment; filename="attendance_report_{month}_{year}.xlsx"'
        )
        return response
